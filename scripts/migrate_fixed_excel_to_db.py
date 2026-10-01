import openpyxl
import sqlite3
import uuid
import re
import json
from datetime import datetime, timezone

excel_path = "/Users/yashbisht/saree-app/DHARMA TEX - REPAIRED.xlsx"
db_path = "/Users/yashbisht/saree-app/prisma/dev.db"

wb = openpyxl.load_workbook(excel_path, data_only=True)
conn = sqlite3.connect(db_path)
cursor = conn.cursor()

def gen_cuid():
    return "c" + uuid.uuid4().hex[:24]

def now_iso():
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")

def dt_to_iso(dt):
    if isinstance(dt, datetime):
        return dt.strftime("%Y-%m-%dT12:00:00.000Z")
    try:
        parsed = datetime.fromisoformat(str(dt))
        return parsed.strftime("%Y-%m-%dT12:00:00.000Z")
    except:
        return now_iso()

def clean_str(s):
    if s is None:
        return ""
    return re.sub(r'\s+', ' ', str(s)).strip()

def norm_key(s):
    return re.sub(r'[^A-Z0-9]', '', clean_str(s).upper())

print("Step 1: Clearing all existing tables in database...")
tables_to_clear = [
    "SaleItem", "Sale", "LedgerEntry", "Customer", "Purchase", "Product", "Supplier", "Expense", "User"
]
for t in tables_to_clear:
    cursor.execute(f"DELETE FROM {t};")
conn.commit()
print("All tables cleared successfully.")

print("\nStep 2: Inserting Suppliers...")
suppliers = {
    "AJMERA TEX": {"id": gen_cuid(), "name": "Ajmera Tex Surat", "phone": "9825100001"},
    "DHARMA TEX": {"id": gen_cuid(), "name": "Dharma Tex Surat", "phone": "9825100002"},
    "M B CREATION": {"id": gen_cuid(), "name": "M B Creation Surat", "phone": "9825100003"},
    "RAJBANNI TEX": {"id": gen_cuid(), "name": "Rajbanni Tex Surat", "phone": "9825100004"},
    "LALA JI TEX": {"id": gen_cuid(), "name": "Lala Ji Tex Surat", "phone": "9825100005"},
    "SUPREME TEX": {"id": gen_cuid(), "name": "Supreme Tex", "phone": "9825100006"},
    "VAISHNAVI IN-HOUSE": {"id": gen_cuid(), "name": "Vaishnavi Collection In-House", "phone": "9876543210"}
}

for s_code, s_info in suppliers.items():
    cursor.execute(
        "INSERT INTO Supplier (id, name, phone) VALUES (?, ?, ?);",
        (s_info["id"], s_info["name"], s_info["phone"])
    )
conn.commit()
print(f"Inserted {len(suppliers)} suppliers.")

print("\nStep 3: Inserting Expenses from EXPENDITURE sheet...")
expenses_data = [
    ("Saree Racks (6 units)", 22800.0, "Fixtures", "Qty: 6, Rate: ₹3,800"),
    ("PVC Curtains (2 units)", 980.0, "Interior", "Qty: 2, Rate: ₹490"),
    ("Carry Bags Custom Printing", 5000.0, "Packaging", "Bulk store packaging bags"),
    ("Store Dressing Mirror", 1900.0, "Fixtures", "Fitting mirror"),
    ("Maintenance & Cleaning Bucket", 400.0, "Maintenance", "Cleaning supplies")
]

for title, amt, cat, note in expenses_data:
    cursor.execute(
        "INSERT INTO Expense (id, title, amount, category, note, date) VALUES (?, ?, ?, ?, ?, ?);",
        (gen_cuid(), title, amt, cat, note, "2025-10-01T10:00:00.000Z")
    )
conn.commit()
print(f"Inserted {len(expenses_data)} store setup expenses (Total: ₹31,080.00).")

print("\nStep 4: Reading SALE PURCHASE stock on hand...")
ws_sp = wb['SALE PURCHASE']
stock_dict = {}
for r in range(4, ws_sp.max_row + 1):
    prod = clean_str(ws_sp.cell(row=r, column=8).value)
    stk = ws_sp.cell(row=r, column=11).value
    price = ws_sp.cell(row=r, column=13).value
    if prod:
        k = norm_key(prod)
        stock_dict[k] = {
            "name": prod,
            "stock": max(0, int(stk or 0)),
            "price": float(price) if price and float(price) > 0 else None
        }

print(f"Loaded {len(stock_dict)} product stock records from SALE PURCHASE.")

print("\nStep 5: Cataloging and inserting Products from Vendor Sheets...")

def infer_cat_fab(name, desc=""):
    n_lower = name.lower()
    d_lower = desc.lower()
    
    if "peticoat" in n_lower or "peticoat" in d_lower:
        return "Petticoat", "Cotton"
    if "blouse" in n_lower or "blouse" in d_lower:
        return "Blouse", "Cotton"
    if "lehenga" in n_lower or "lehenge" in d_lower:
        return "Lehenga", "Net" if "net" in n_lower else "Silk"
    if any(k in n_lower for k in ["kurti", "suit", "cts"]):
        return "Suit", "Cotton" if "cotton" in n_lower else "Silk"
    if "georgette" in n_lower:
        return "Saree", "Georgette"
    if "bandhani" in n_lower:
        return "Saree", "Cotton"
    if "dobby" in n_lower:
        return "Saree", "Dobby Silk"
    if "chiffon" in n_lower or "siffon" in n_lower:
        return "Saree", "Chiffon"
    if "cotton" in n_lower:
        return "Saree", "Cotton"
    if "silk" in n_lower:
        return "Saree", "Silk"
    return "Saree", "Silk"

products_map = {}
sku_counter = 1
vendor_purchases_items = {k: [] for k in suppliers.keys()}

def add_product(name, supplier_key, category, fabric, cost_price, selling_price, initial_stock, location="Main Rack"):
    global sku_counter
    k = norm_key(name)
    if k in products_map:
        return products_map[k]
    
    sku = f"VS-{category[:3].upper()}-{sku_counter:04d}"
    sku_counter += 1
    p_id = gen_cuid()
    
    final_stock = initial_stock
    if k in stock_dict:
        final_stock = stock_dict[k]["stock"]
        if stock_dict[k]["price"]:
            selling_price = max(selling_price, stock_dict[k]["price"])
            
    cost_price = round(float(cost_price), 2)
    selling_price = round(float(selling_price), 2)
    if selling_price <= cost_price:
        selling_price = round(cost_price * 1.4, 2)
        
    p_data = {
        "id": p_id,
        "sku": sku,
        "name": clean_str(name),
        "category": category,
        "fabric": fabric,
        "costPrice": cost_price,
        "sellingPrice": selling_price,
        "stockQty": int(final_stock),
        "minStock": 5,
        "location": location,
        "supplierId": suppliers[supplier_key]["id"],
        "createdAt": now_iso(),
        "updatedAt": now_iso()
    }
    
    cursor.execute(
        """INSERT INTO Product (id, sku, name, category, fabric, costPrice, sellingPrice, stockQty, minStock, location, supplierId, createdAt, updatedAt)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);""",
        (p_data["id"], p_data["sku"], p_data["name"], p_data["category"], p_data["fabric"],
         p_data["costPrice"], p_data["sellingPrice"], p_data["stockQty"], p_data["minStock"],
         p_data["location"], p_data["supplierId"], p_data["createdAt"], p_data["updatedAt"])
    )
    products_map[k] = p_data
    
    # Track for purchase bill
    vendor_purchases_items[supplier_key].append({
        "productId": p_id,
        "name": p_data["name"],
        "qty": int(initial_stock),
        "rate": cost_price,
        "total": round(initial_stock * cost_price, 2)
    })
    
    return p_data

# 1. AJMERA TEX
ws = wb['AJMERA TEX']
for r in range(3, ws.max_row + 1):
    item = ws.cell(row=r, column=2).value
    if not item: continue
    desc = clean_str(ws.cell(row=r, column=4).value)
    qty = int(ws.cell(row=r, column=5).value or 0)
    rate = float(ws.cell(row=r, column=6).value or 0)
    cost = float(ws.cell(row=r, column=8).value or rate)
    sp = float(ws.cell(row=r, column=15).value or ws.cell(row=r, column=14).value or (rate * 1.5))
    cat, fab = infer_cat_fab(str(item), desc)
    loc = "Petticoat Section" if cat == "Petticoat" else "Blouse Counter" if cat == "Blouse" else "Rack A"
    add_product(item, "AJMERA TEX", cat, fab, cost, sp, qty, loc)

# 2. D-TEX
ws = wb['D-TEX']
for r in range(3, ws.max_row + 1):
    item = ws.cell(row=r, column=2).value
    if not item: continue
    qty = int(ws.cell(row=r, column=3).value or 0)
    rate = float(ws.cell(row=r, column=4).value or 0)
    cost = float(ws.cell(row=r, column=7).value or rate)
    sp = float(ws.cell(row=r, column=13).value or (rate * 1.5))
    cat, fab = infer_cat_fab(str(item))
    add_product(item, "DHARMA TEX", cat, fab, cost, sp, qty, "Rack B")

# 3. LEHENGE-M B Creation
ws = wb['LEHENGE-M B Creation']
for r in range(3, ws.max_row + 1):
    item = ws.cell(row=r, column=2).value
    if not item: continue
    qty = int(ws.cell(row=r, column=3).value or 0)
    rate = float(ws.cell(row=r, column=4).value or 0)
    cost = float(ws.cell(row=r, column=7).value or rate)
    sp = float(ws.cell(row=r, column=11).value or (rate * 1.5))
    cat, fab = infer_cat_fab(str(item), "Lehenga")
    add_product(item, "M B CREATION", "Lehenga", fab, cost, sp, qty, "Lehenga Wardrobe")

# 4. RAJBANNI TEX
ws = wb['RAJBANNI TEX']
for r in range(3, ws.max_row + 1):
    item = ws.cell(row=r, column=2).value
    if not item: continue
    qty = int(ws.cell(row=r, column=3).value or 0)
    rate = float(ws.cell(row=r, column=4).value or 0)
    cost = float(ws.cell(row=r, column=7).value or rate)
    sp = float(ws.cell(row=r, column=13).value or ws.cell(row=r, column=11).value or (rate * 1.5))
    cat, fab = infer_cat_fab(str(item))
    add_product(item, "RAJBANNI TEX", cat, fab, cost, sp, qty, "Rack C")

# 5. LALA JI TEX
ws = wb['LALA JI TEX']
for r in range(3, ws.max_row + 1):
    item = ws.cell(row=r, column=2).value
    if not item: continue
    qty = int(ws.cell(row=r, column=3).value or 0)
    rate = float(ws.cell(row=r, column=4).value or 0)
    cost = float(ws.cell(row=r, column=7).value or rate)
    sp = float(ws.cell(row=r, column=13).value or ws.cell(row=r, column=11).value or (rate * 1.5))
    cat, fab = infer_cat_fab(str(item))
    add_product(item, "LALA JI TEX", cat, fab, cost, sp, qty, "Rack D")

# 6. suits
ws = wb['suits']
for r in range(3, ws.max_row + 1):
    item = ws.cell(row=r, column=2).value
    if not item: continue
    qty = int(ws.cell(row=r, column=3).value or 0)
    rate = float(ws.cell(row=r, column=4).value or 0)
    cost = float(ws.cell(row=r, column=7).value or rate)
    sp = float(ws.cell(row=r, column=12).value or (rate * 1.5))
    cat, fab = infer_cat_fab(str(item), "Suit")
    add_product(item, "SUPREME TEX", "Suit", fab, cost, sp, qty, "Suit Section")

# Also add any product listed in SALE PURCHASE stock summary that wasn't in vendor sheets
for k, s_info in stock_dict.items():
    if k not in products_map:
        cat, fab = infer_cat_fab(s_info["name"])
        price = s_info["price"] or 1000.0
        cost = round(price * 0.65, 2)
        add_product(s_info["name"], "VAISHNAVI IN-HOUSE", cat, fab, cost, price, s_info["stock"], "Store Floor")

conn.commit()
print(f"Total Products inserted into database: {len(products_map)}")

print("\nStep 6: Inserting Vendor Invoices into Purchase table...")
for s_code, items in vendor_purchases_items.items():
    if items:
        tot_cost = sum(it["total"] for it in items)
        bill_no = f"PO-{s_code[:3]}-2025"
        cursor.execute(
            """INSERT INTO Purchase (id, supplierId, billNo, items, totalCost, createdAt)
               VALUES (?, ?, ?, ?, ?, ?);""",
            (gen_cuid(), suppliers[s_code]["id"], bill_no, json.dumps(items), tot_cost, "2025-10-06T10:00:00.000Z")
        )
conn.commit()
print("Inserted vendor purchase orders.")

print("\nStep 7: Inserting Default Walk-in Customer...")
default_cust_id = gen_cuid()
cursor.execute(
    """INSERT INTO Customer (id, name, phone, address, balanceDue, totalSpent, createdAt)
       VALUES (?, ?, ?, ?, ?, ?, ?);""",
    (default_cust_id, "Walk-in Store Customer", "9999900000", "Almora Market", 0.0, 0.0, now_iso())
)
conn.commit()

print("\nStep 8: Inserting Historical Sales from Daily Sheet...")
ws_daily = wb['Daily Sheet']
sales_count = 0
items_count = 0
total_sales_val = 0.0

def find_or_create_product(name, fallback_price):
    k = norm_key(name)
    if k in products_map:
        return products_map[k]
    
    for existing_k, p in products_map.items():
        if len(existing_k) > 4 and (existing_k in k or k in existing_k):
            return p
            
    cat, fab = infer_cat_fab(name)
    sp = float(fallback_price) if fallback_price and float(fallback_price) > 0 else 500.0
    cost = round(sp * 0.65, 2)
    return add_product(name, "VAISHNAVI IN-HOUSE", cat, fab, cost, sp, 10, "Counter Stock")

bill_seq = 1001

for r in range(3, 442):
    p_name = clean_str(ws_daily.cell(row=r, column=3).value)
    if not p_name: continue
    
    dt = ws_daily.cell(row=r, column=2).value
    qty = ws_daily.cell(row=r, column=4).value or 1
    price = ws_daily.cell(row=r, column=5).value or 0
    tot = ws_daily.cell(row=r, column=6).value
    
    qty = int(qty)
    price = float(price)
    tot = float(tot) if tot is not None else (qty * price)
    
    prod_record = find_or_create_product(p_name, price)
    
    created_at = dt_to_iso(dt)
    bill_no = f"VS-{created_at[:10].replace('-', '')}-{bill_seq}"
    bill_seq += 1
    sale_id = gen_cuid()
    
    cursor.execute(
        """INSERT INTO Sale (id, billNo, customerId, subtotal, discount, total, paymentMode, amountPaid, amountDue, createdAt)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);""",
        (sale_id, bill_no, default_cust_id, tot, 0.0, tot, "Cash" if tot < 2000 else "UPI", tot, 0.0, created_at)
    )
    
    sale_item_id = gen_cuid()
    cursor.execute(
        """INSERT INTO SaleItem (id, saleId, productId, qty, price, total)
           VALUES (?, ?, ?, ?, ?, ?);""",
        (sale_item_id, sale_id, prod_record["id"], qty, price, tot)
    )
    
    sales_count += 1
    items_count += 1
    total_sales_val += tot

cursor.execute("UPDATE Customer SET totalSpent = ? WHERE id = ?;", (total_sales_val, default_cust_id))

print(f"Inserted {sales_count} sales transactions with {items_count} line items.")
print(f"Total historical sales revenue loaded: ₹{total_sales_val:,.2f}")

print("\nStep 9: Inserting Store Owner User...")
owner_id = gen_cuid()
owner_hash = "$2b$10$.pW8q35KLKTxPznD9DPjbuivJa//CDV0Vee68BGrVS1hQs9y7tusW"
cursor.execute(
    """INSERT INTO User (id, name, username, password, role)
       VALUES (?, ?, ?, ?, ?);""",
    (owner_id, "Store Owner", "owner", owner_hash, "ADMIN")
)
conn.commit()
print("Store Owner user inserted successfully (username: owner).")

conn.close()
print("\n=== DATABASE MIGRATION COMPLETE ===")
